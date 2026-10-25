import io
import json
import zipfile
from datetime import timedelta

from django.conf import settings
from django.db import transaction
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from django.views.decorators.http import require_POST, require_GET

from core.payloads import SCHEMA, build_payload, normalize_content
from core.qr import appearance_settings, render_export
from .access import consume, private
from .imports import read_workbook
from .models import ImportBatch, QRRecord, BatchOutput, public_id
from .views import api, body


@private
@require_GET
@api
def template(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    kind = request.GET.get("kind", "vcard")
    if kind not in ("vcard", "website"):
        raise ValueError("Choose vCard or website.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    keys = ["name", *[k for k, _, _ in SCHEMA[kind]]]
    sample = {
        "name": "Example QR",
        "first": "Alex",
        "last": "Example",
        "organization": "Example Studio",
        "email": "alex@example.com",
        "phone": "+12025550100",
        "website": "https://example.com",
        "url": "https://example.com",
    }
    sheet.append(keys)
    sheet.append([sample.get(k, "") for k in keys])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, key in enumerate(keys, 1):
        cell = sheet.cell(1, index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3685")
        sheet.column_dimensions[get_column_letter(index)].width = 24
        if key in ("phone", "office", "direct", "postal"):
            for row in range(2, 502):
                sheet.cell(row, index).number_format = "@"
    instructions = workbook.create_sheet("Instructions")
    for message in [
        "Fill in the Data sheet. Replace or delete the example row.",
        "Keep header names unchanged for automatic column mapping.",
        "Maximum 500 data rows; do not use formulas or error values.",
        "Store phone numbers and postal codes as TEXT to preserve leading zeros.",
        "Website addresses must begin with https:// or http://.",
        "Import creates new QR codes; it never updates existing ones.",
        "For vCard custom fields add a column named custom:Department (replace Department).",
    ]:
        instructions.append([message])
    instructions.column_dimensions["A"].width = 100
    stream = io.BytesIO()
    workbook.save(stream)
    response = HttpResponse(
        stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="qr-{kind}-template.xlsx"'
    return response


def batch(id, locking=False):
    qs = ImportBatch.objects.select_for_update() if locking else ImportBatch.objects
    return get_object_or_404(qs, pk=id, expires_at__gt=timezone.now())


def report(b):
    return {
        "id": b.id,
        "status": b.status,
        "cursor": b.cursor,
        "total": len(b.rows),
        "valid": sum(not r.get("error") for r in b.rows),
        "invalid": sum(bool(r.get("error")) for r in b.rows),
        "completed": sum(bool(r.get("completed")) for r in b.rows),
        "rows": [{k: v for k, v in r.items() if k != "data"} for r in b.rows],
    }


@private
@require_POST
@api
def inspect(request):
    if not consume(request, "import", 20, 3600):
        return JsonResponse({"error": "Import limit reached. Try later."}, status=429)
    upload = request.FILES.get("file")
    if upload is None or not upload.name.lower().endswith(".xlsx") or upload.size > 5 * 1024 * 1024:
        raise ValueError("Choose an .xlsx file no larger than 5 MiB.")
    sheets = read_workbook(upload.read())
    b = ImportBatch.objects.create(sheets=sheets, expires_at=timezone.now() + timedelta(hours=24))
    return JsonResponse(
        {
            "id": b.id,
            "sheets": {
                name: {
                    "headers": [v["value"] for v in rows[0]] if rows else [],
                    "preview": [[v["value"] for v in row] for row in rows[1:6]],
                    "rows": max(0, len(rows) - 1),
                }
                for name, rows in sheets.items()
            },
        }
    )


@private
@require_POST
@api
def validate(request):
    value = body(request)
    kind, mode = value.get("kind"), value.get("mode")
    if kind not in ("vcard", "website") or mode not in ("static", "persistent"):
        raise ValueError("Choose vCard or website, static or persistent.")
    with transaction.atomic():
        b = batch(value.get("id"), True)
        if b.status not in ("uploaded", "validated"):
            raise ValueError("This import has already started. Upload again for a new import.")
        sheet = b.sheets.get(value.get("sheet"))
        if not sheet:
            raise ValueError("Choose a nonempty worksheet.")
        mapping = value.get("mapping")
        allowed = {k for k, _, _ in SCHEMA[kind]} | {"name", ""}
        if (
            not isinstance(mapping, list)
            or len(mapping) != len(sheet[0])
            or any(
                not isinstance(m, str)
                or (m not in allowed and not (kind == "vcard" and m.startswith("custom:")))
                for m in mapping
            )
        ):
            raise ValueError("Map each column to a supported field or Ignore.")
        used = [v for v in mapping if v]
        if not used or len(used) != len(set(used)):
            raise ValueError("Choose at least one field. Each field may be mapped only once.")
        b.appearance = appearance_settings(value.get("appearance", {}))
        output = []
        for number, cells in enumerate(sheet[1:], 2):
            if not any(c["value"] for c in cells):
                continue
            row = {"row": number, "warnings": [], "completed": False}
            data, custom, error = {}, [], ""
            for target, cell in zip(mapping, cells):
                if not target:
                    continue
                if cell["type"] in ("f", "e"):
                    error = "Formula or error cell: replace it with a literal value."
                if target in ("phone", "office", "direct") and cell["type"] == "n" and cell["value"]:
                    row["warnings"].append(
                        "Numeric phone cell: leading zeros may already be lost. Store phones as text."
                    )
                if target.startswith("custom:"):
                    if cell["value"]:
                        custom.append({"name": target[7:], "value": cell["value"]})
                else:
                    data[target] = cell["value"]
            data["custom"] = custom
            try:
                if error:
                    raise ValueError(error)
                row["data"] = normalize_content(kind, data)
                row["name"] = (
                    data.get("name")
                    or " ".join(filter(None, (data.get("first"), data.get("last"))))
                    or data.get("organization")
                    or data.get("url")
                    or data.get("email")
                    or f"Row {number}"
                )[:120]
            except ValueError as exc:
                row["error"] = str(exc)
            output.append(row)
        b.kind, b.mode, b.rows, b.status = kind, mode, output, "validated"
        b.save()
    return JsonResponse(report(b))


@private
@require_POST
@api
def step(request, id):
    expected = body(request).get("cursor")
    with transaction.atomic():
        b = batch(id, True)
        if b.status in ("cancelled", "complete") or expected != b.cursor:
            return JsonResponse(report(b))
        if b.status not in ("validated", "running"):
            raise ValueError("Validate the import first.")
        b.status = "running"
        end = min(b.cursor + 5, len(b.rows))
        for row in b.rows[b.cursor : end]:
            if row.get("error"):
                continue
            try:
                id = public_id()
                payload = (
                    settings.PUBLIC_BASE_URL + "/q/" + id + "/"
                    if b.mode == "persistent"
                    else build_payload(b.kind, row["data"])
                )
                png = render_export(payload, b.appearance, "png")
                svg = render_export(payload, b.appearance, "svg")
                size = len(png) + len(svg)
                if b.output_bytes + size > 64 * 1024 * 1024:
                    raise ValueError("Batch output exceeds 64 MiB; split the import or shorten the content.")
                if b.mode == "persistent":
                    record = QRRecord.objects.create(
                        id=id, name=row["name"], kind=b.kind, data=row["data"], appearance=b.appearance
                    )
                    row["id"] = record.id
                BatchOutput.objects.create(batch=b, row=row["row"], png=png, svg=svg)
                b.output_bytes += size
                row["completed"] = True
            except ValueError as exc:
                row["error"] = str(exc)
        b.cursor = end
        if end == len(b.rows):
            b.status = "complete"
            b.sheets = {}
        b.save()
    return JsonResponse(report(b))


@private
@require_POST
def cancel(request, id):
    with transaction.atomic():
        b = batch(id, True)
        b.status = "cancelled"
        b.sheets = {}
        b.save(update_fields=["status", "sheets"])
    return JsonResponse(report(b))


@private
@require_GET
@api
def download(request, id):
    b = batch(id)
    fmt = request.GET.get("format", "png")
    if fmt not in ("png", "svg"):
        raise ValueError("Choose PNG or SVG.")
    if not consume(request, "zip", 20, 3600):
        return JsonResponse({"error": "Download limit reached. Try later."}, status=429)
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("report.json", json.dumps(report(b), ensure_ascii=False, indent=2))
        for output in BatchOutput.objects.filter(batch=b).only("row", fmt).iterator(chunk_size=5):
            archive.writestr(f"row-{output.row:04d}.{fmt}", bytes(getattr(output, fmt)))
    response = HttpResponse(stream.getvalue(), content_type="application/zip")
    response["Content-Disposition"] = 'attachment; filename="qr-export.zip"'
    return response


@private
@require_GET
def status(request, id):
    return JsonResponse(report(batch(id)))
