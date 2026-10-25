"""Bounded workbook parsing; never evaluate formulas or expose original uploads."""

import io
import zipfile
from contextlib import closing

from openpyxl import load_workbook


def read_workbook(content):
    if len(content) > 5 * 1024 * 1024:
        raise ValueError("Workbook exceeds 5 MiB.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 500
                or sum(f.file_size for f in entries) > 40 * 1024 * 1024
                or any(f.flag_bits & 1 for f in entries)
            ):
                raise ValueError("Expanded workbook exceeds the safety limits.")
        with closing(
            load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        ) as workbook:
            if len(workbook.worksheets) > 20:
                raise ValueError("Use at most 20 worksheets.")
            sheets = {}
            cells_total = 0
            for sheet in workbook.worksheets:
                if (sheet.max_row or 0) > 501 or (sheet.max_column or 0) > 100:
                    raise ValueError("Use at most 500 data rows and 100 columns per sheet.")
                rows = []
                for row in sheet.iter_rows():
                    converted = []
                    for cell in row:
                        cells_total += 1
                        if cells_total > 50000:
                            raise ValueError("Use at most 50,000 cells per workbook.")
                        value = "" if cell.value is None else str(cell.value)
                        if len(value) > 4000:
                            raise ValueError("A cell exceeds 4,000 characters.")
                        converted.append({"value": value, "type": cell.data_type})
                    rows.append(converted)
                sheets[sheet.title] = rows
            return sheets
    except ValueError:
        raise
    except Exception:
        # XLSX is an untrusted ZIP/XML boundary; parser messages can include user data.
        raise ValueError("This is not a valid, supported .xlsx workbook.") from None
